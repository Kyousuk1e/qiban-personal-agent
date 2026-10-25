"""Bounded document text extraction, executed in a separate local process."""

import json
import sys
import zipfile
from pathlib import Path

TEXT_TYPES = {
    ".txt",
    ".md",
    ".csv",
    ".json",
    ".py",
    ".toml",
    ".yaml",
    ".yml",
    ".ini",
    ".log",
}
DOCUMENT_TYPES = {".pdf", ".docx", ".xlsx", ".pptx"}
SUPPORTED = TEXT_TYPES | DOCUMENT_TYPES
MAX_UPLOAD = 20 * 1024 * 1024
MAX_TEXT = 180000


def extract(path):
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED:
        raise ValueError("不支持该格式。旧版 doc/xls/ppt 请先另存为 docx/xlsx/pptx。")
    if path.stat().st_size > MAX_UPLOAD:
        raise ValueError("原文件不能超过 20 MB。")
    warnings = []
    pieces, size, clipped = [], 0, False

    def add(value):
        nonlocal size, clipped
        value = str(value).strip()
        if not value:
            return True
        raw = (value + "\n\n").encode("utf-8")
        remaining = MAX_TEXT - size
        if len(raw) > remaining:
            pieces.append(raw[:remaining].decode("utf-8", errors="ignore"))
            size = MAX_TEXT
            clipped = True
            return False
        pieces.append(raw.decode("utf-8"))
        size += len(raw)
        return True

    if suffix in TEXT_TYPES:
        raw = path.read_bytes()
        try:
            content = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            try:
                content = raw.decode("gb18030")
                warnings.append("原文件按 GB18030 解码，已转为 UTF-8。")
            except UnicodeDecodeError:
                raise ValueError("无法识别文本编码，请另存为 UTF-8 后重试。") from None
        if any(ord(c) < 32 and c not in "\t\n\r" for c in content):
            raise ValueError("文件包含二进制内容，不能当作文本导入。")
        add(content)
    elif suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(path)
        if reader.is_encrypted and not reader.decrypt(""):
            raise ValueError("PDF 有密码保护，请先解密后上传。")
        if len(reader.pages) > 100:
            warnings.append("仅提取前 100 页，后续页未读取。")
        missing = []
        for index, page in enumerate(reader.pages[:100], 1):
            content = page.extract_text() or ""
            if content.strip():
                if not add(f"## 第 {index} 页\n{content}"):
                    break
            else:
                missing.append(str(index))
        if missing:
            warnings.append(
                "第 "
                + "、".join(missing)
                + " 页无可提取文字，可能为扫描页；尚未进行 OCR。"
            )
    else:
        # Office XML containers are never unpacked to user-controlled paths.
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            if (
                len(entries) > 3000
                or sum(e.file_size for e in entries) > 40 * 1024 * 1024
            ):
                raise ValueError("Office 文件解压后体积或条目数过大，无法安全解析。")
            if any("vbaproject" in e.filename.lower() for e in entries):
                raise ValueError("不接受包含宏的 Office 文件。")
        if suffix == ".docx":
            from docx import Document
            from docx.table import Table
            from docx.text.paragraph import Paragraph

            doc = Document(path)
            for element in doc.iter_inner_content():
                if isinstance(element, Paragraph):
                    content = element.text
                elif isinstance(element, Table):
                    content = "\n".join(
                        " | ".join(c.text for c in r.cells) for r in element.rows
                    )
                else:
                    continue
                if not add(content):
                    break
            warnings.append("已提取正文和表格；图片、文本框、页眉页脚等未解析。")
        elif suffix == ".xlsx":
            from openpyxl import load_workbook

            workbook = load_workbook(
                path, read_only=True, data_only=False, keep_links=False
            )
            try:
                for sheet in workbook.worksheets[:20]:
                    if not add(f"## 工作表：{sheet.title}"):
                        break
                    for index, row in enumerate(
                        sheet.iter_rows(max_row=1000, max_col=40, values_only=True), 1
                    ):
                        if any(value is not None for value in row):
                            if not add(
                                f"第 {index} 行："
                                + " | ".join("" if v is None else str(v) for v in row)
                            ):
                                break
                    if clipped:
                        break
            finally:
                workbook.close()
            warnings.append(
                "最多读取前 20 个工作表，每表前 1000 行、40 列；公式仅保留原文，不计算，图表与图片未解析。"
            )
        elif suffix == ".pptx":
            from pptx import Presentation

            presentation = Presentation(path)
            for index, slide in enumerate(list(presentation.slides)[:100], 1):
                content = []
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        content.append(shape.text)
                    if shape.has_table:
                        content.extend(
                            " | ".join(c.text for c in r.cells)
                            for r in shape.table.rows
                        )
                if content and not add(
                    f"## 第 {index} 页幻灯片\n" + "\n".join(content)
                ):
                    break
            warnings.append(
                "最多读取前 100 张幻灯片的文本与表格；图片、图表、组合对象及备注未解析。"
            )
    if clipped:
        warnings.append("提取文本超过 180 KB，已截断；只可根据已提取内容回答。")
    content = "".join(pieces).strip()
    if not content:
        warnings.append(
            "未提取到可供助手阅读的文字，不能据此总结原文件；扫描文档需先 OCR。"
        )
    return {
        "text": content,
        "warning": " ".join(warnings),
        "status": "ready" if content else "empty",
    }


if __name__ == "__main__":
    try:
        result = extract(sys.argv[1])
    except Exception as error:
        result = {
            "error": str(error)
            if isinstance(error, ValueError)
            else "文件损坏或格式不匹配，无法提取文字。"
        }
    Path(sys.argv[2]).write_text(
        json.dumps(result, ensure_ascii=False), encoding="utf-8"
    )
