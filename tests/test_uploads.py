import io
import re
import subprocess
import tempfile
import unittest
from unittest.mock import patch

from fastapi.testclient import TestClient

from qiban.agent import execute_tool
from qiban.extract import MAX_UPLOAD
from qiban.files import Files
from qiban.storage import Store
from qiban.uploads import Uploads
from qiban.web import create_app


class UploadTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.store = Store(self.tmp.name)
        self.uploads = Uploads(self.store)
        self.client = TestClient(
            create_app(self.store, start_worker=False), base_url="http://127.0.0.1"
        )
        page = self.client.get("/files")
        self.csrf = re.search(r'name="_csrf" value="([^"]+)"', page.text)[1]

    def tearDown(self):
        self.client.close()
        self.tmp.cleanup()

    def upload(self, name, content):
        return self.client.post(
            "/files/upload",
            data={"_csrf": self.csrf},
            files={"file": (name, content, "application/octet-stream")},
        )

    def test_browser_upload_preserves_original_and_duplicate_names(self):
        raw = "# 项目需求\n主要使用 Python。".encode("utf-8")
        for _ in range(2):
            response = self.upload("需求.md", raw)
            self.assertEqual(response.status_code, 200, response.text)
            self.assertIn("已上传：需求.md", response.text)
        items = self.uploads.all()
        self.assertEqual(len(items), 2)
        self.assertNotEqual(items[0]["text_name"], items[1]["text_name"])
        for item in items:
            self.assertEqual(
                self.client.get("/files/original/" + item["id"]).content, raw
            )
            self.assertIn(
                "主要使用 Python", Files(self.store.root).read(item["text_name"])
            )
        response = self.client.post(
            "/files/summarize/" + items[0]["id"], data={"_csrf": self.csrf}
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn(
            items[0]["text_name"], self.store.one("SELECT prompt FROM tasks")["prompt"]
        )

    def test_missing_csrf_and_bad_files(self):
        response = self.client.post("/files/upload", files={"file": ("x.txt", b"abc")})
        self.assertEqual(response.status_code, 403)
        for name, content in [
            ("../bad.txt", b"a"),
            ("x.exe", b"a"),
            ("x.doc", b"a"),
            ("x.txt", b"\x00\x01"),
            ("x.pdf", b"not a PDF"),
        ]:
            with self.subTest(name=name):
                self.assertEqual(self.upload(name, content).status_code, 400)
        self.assertEqual(self.uploads.all(), [])
        self.assertEqual(list(self.uploads.root.iterdir()), [])

    def test_upload_size_limit(self):
        response = self.upload("huge.txt", b"x" * (MAX_UPLOAD + 70000))
        self.assertEqual(response.status_code, 413)
        self.assertEqual(self.uploads.all(), [])

    def test_docx_xlsx_pptx_text_extraction(self):
        from docx import Document
        from openpyxl import Workbook
        from pptx import Presentation
        from pptx.util import Inches

        doc = Document()
        doc.add_paragraph("Document fixture")
        table = doc.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "Table fixture"
        content = io.BytesIO()
        doc.save(content)
        item = self.uploads.import_file("notes.docx", content.getvalue())
        self.assertIn("Table fixture", Files(self.store.root).read(item["text_name"]))
        workbook = Workbook()
        workbook.active["A1"] = "Sheet fixture"
        workbook.active["B1"] = "=1+2"
        content = io.BytesIO()
        workbook.save(content)
        item = self.uploads.import_file("plan.xlsx", content.getvalue())
        self.assertIn("=1+2", Files(self.store.root).read(item["text_name"]))
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(
            Inches(1), Inches(1), Inches(4), Inches(1)
        ).text = "Slide fixture"
        content = io.BytesIO()
        deck.save(content)
        item = self.uploads.import_file("slides.pptx", content.getvalue())
        self.assertIn("Slide fixture", Files(self.store.root).read(item["text_name"]))

    def test_pdf_text_and_scanned_page_are_distinguished(self):
        from pypdf import PdfWriter
        from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

        writer = PdfWriter()
        page = writer.add_blank_page(width=400, height=400)
        page[NameObject("/Resources")] = DictionaryObject(
            {
                NameObject("/Font"): DictionaryObject(
                    {
                        NameObject("/F1"): DictionaryObject(
                            {
                                NameObject("/Type"): NameObject("/Font"),
                                NameObject("/Subtype"): NameObject("/Type1"),
                                NameObject("/BaseFont"): NameObject("/Helvetica"),
                            }
                        )
                    }
                )
            }
        )
        stream = DecodedStreamObject()
        stream.set_data(b"BT /F1 12 Tf 20 300 Td (PDF fixture) Tj ET")
        page[NameObject("/Contents")] = stream
        content = io.BytesIO()
        writer.write(content)
        item = self.uploads.import_file("guide.pdf", content.getvalue())
        self.assertEqual(item["status"], "ready")
        self.assertIn("PDF fixture", Files(self.store.root).read(item["text_name"]))
        writer = PdfWriter()
        writer.add_blank_page(width=400, height=400)
        content = io.BytesIO()
        writer.write(content)
        item = self.uploads.import_file("scan.pdf", content.getvalue())
        self.assertEqual(item["status"], "empty")
        self.assertIn("OCR", item["warning"])
        self.assertEqual(
            self.client.post(
                "/files/summarize/" + item["id"], data={"_csrf": self.csrf}
            ).status_code,
            400,
        )

    def test_long_text_pagination_and_truncation_notice(self):
        item = self.uploads.import_file("long.txt", ("测试资料" * 20000).encode())
        self.assertIn("截断", item["warning"])
        full = Files(self.store.root).read(item["text_name"])
        first = execute_tool(self.store, {}, "read_file", {"name": item["text_name"]})
        second = execute_tool(
            self.store,
            {},
            "read_file",
            {"name": item["text_name"], "offset": first["next_offset"]},
        )
        self.assertEqual(first["text"] + second["text"], full[:24000])
        catalog = execute_tool(self.store, {}, "list_files", {})
        self.assertEqual(catalog["uploads"][0]["name"], "long.txt")

    def test_timeout_cleans_staged_original(self):
        with patch(
            "qiban.uploads.subprocess.run",
            side_effect=subprocess.TimeoutExpired("parser", 30),
        ):
            with self.assertRaisesRegex(ValueError, "30"):
                self.uploads.import_file("timeout.txt", b"test")
        self.assertEqual(list(self.uploads.root.iterdir()), [])


if __name__ == "__main__":
    unittest.main()
